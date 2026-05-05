import docx
from docx.shared import Pt

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            doc.add_heading(line[2:], level=0)
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            run = p.add_run(line[5:])
            run.bold = True
        elif line.startswith('- [ ] ') or line.startswith('- '):
            p = doc.add_paragraph(line, style='List Bullet')
        elif line == '---':
            doc.add_page_break()
        else:
            doc.add_paragraph(line)
            
    doc.save(docx_path)
    print(f"Successfully saved to {docx_path}")

if __name__ == "__main__":
    md_file = r"C:\Users\Acer\.gemini\antigravity\brain\7d344070-0948-4c60-a3b2-6710077e24cc\ai_batch_regrade_test_cases.md"
    output_file = r"d:\OEM_Project\AI_Batch_Regrade_Tool_Test_Cases.docx"
    convert_md_to_docx(md_file, output_file)
